import urllib.request
import json
import os
from docx import Document

url = 'https://smart-edudigital.github.io/smartdrill-sd-2026/questions.json'
req = urllib.request.urlopen(url)
data = json.load(req)

target_dir = r"C:/Users/MIP-USER/Documents/Project Soal/web_pulled_docx"
os.makedirs(target_dir, exist_ok=True)

# Group by subject and paket
grouped = {}
for q in data:
    subj = q.get('subject', 'Umum')
    paket = q.get('paket', 'Paket 1')
    key = (subj, paket)
    if key not in grouped:
        grouped[key] = []
    grouped[key].append(q)

for (subj, paket), q_list in grouped.items():
    doc = Document()
    doc.add_heading(f"Soal Try Out - {subj} ({paket})", 0)
    doc.add_paragraph(f"Sumber: Web Try Out (GitHub Pages) | Total Soal: {len(q_list)}")
    
    for i, q in enumerate(q_list, 1):
        p = doc.add_paragraph()
        p.add_run(f"Soal {i} ({q.get('tipe', 'Pilihan Ganda')})\n").bold = True
        p.add_run(f"Materi: {q.get('materi', '-')}\n")
        p.add_run(f"Pertanyaan:\n{q.get('question', '')}\n\n")
        
        opts = q.get('options', [])
        if isinstance(opts, list):
            for idx, opt in enumerate(opts):
                opt_letter = chr(65 + idx)
                p.add_run(f"  {opt_letter}. {opt}\n")
        elif isinstance(opts, dict):
            for k, v in opts.items():
                p.add_run(f"  {k}. {v}\n")
                
        p.add_run(f"\nKunci Jawaban: {q.get('correct', '-')}\n")
        p.add_run(f"Pembahasan: {q.get('explanation', '-')}\n")
        doc.add_paragraph("-" * 40)
        
    filename = f"TryOut_{subj.replace(' ', '_')}_{paket.replace(' ', '_')}.docx"
    filepath = os.path.join(target_dir, filename)
    doc.save(filepath)
    print(f"Saved: {filepath}")

print("Done pulling and generating docx files!")
