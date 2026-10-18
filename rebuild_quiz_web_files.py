import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

TARGET_DIR = r"C:\Users\MIP-USER\Documents\Project Soal"

with open(os.path.join(TARGET_DIR, "questions.json"), encoding="utf-8") as f:
    all_questions = json.load(f)

# Group by subject and paket
grouped = {}
for q in all_questions:
    key = (q["subject"], q["paket"])
    if key not in grouped:
        grouped[key] = []
    grouped[key].append(q)

def buat_docx(paket_data, subject_name, paket_name, filename):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_title.add_run("ASESMEN STANDARISASI PENDIDIKAN DAERAH / TKA SD 2026\n")
    r_sub.font.size = Pt(10)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    
    r_main = p_title.add_run(f"TRY OUT DIGITAL INTERAKTIF: QUIZ WEB\n{subject_name.upper()} - {paket_name.upper()}\n")
    r_main.font.size = Pt(15)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r_desc = p_title.add_run("Disusun Berdasarkan Kisi-Kisi Kurikulum Merdeka & Standar Try Out Resmi Pemerintah")
    r_desc.font.size = Pt(9.5)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.add_run("―" * 55).font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    
    p_petunjuk = doc.add_paragraph()
    p_petunjuk.paragraph_format.space_after = Pt(8)
    p_petunjuk.add_run("PETUNJUK UMUM:\n").font.bold = True
    p_petunjuk.add_run("1. Tulislah nama dan identitas Anda pada lembar pengerjaan yang tersedia.\n")
    p_petunjuk.add_run("2. Periksa dan bacalah soal-soal secara teliti sebelum menjawab.\n")
    p_petunjuk.add_run("3. Naskah soal ini terdiri atas Pilihan Ganda Sederhana, Pilihan Ganda Kompleks, dan Tabel Benar/Salah.\n")
    p_petunjuk.add_run("4. Pilihlah jawaban yang paling benar dan tepat.")
    
    h_sec_a = doc.add_heading(level=2)
    h_sec_a.paragraph_format.space_before = Pt(12)
    h_sec_a.paragraph_format.space_after = Pt(8)
    r_a = h_sec_a.add_run("BAGIAN A: NASKAH SOAL UJIAN (20 BUTIR)")
    r_a.font.size = Pt(12)
    r_a.font.bold = True
    r_a.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    for item in paket_data:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        
        run_num = p_q.add_run(f"{item['number']}. ")
        run_num.font.bold = True
        p_q.add_run(item["question"])
        
        tipe = item.get("tipe", "Pilihan Ganda")
        if tipe == "MCMA":
            for opt_idx, opt_text in enumerate(item["options"]):
                opt_letter = chr(65 + opt_idx)
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Inches(0.3)
                p_opt.paragraph_format.space_before = Pt(0)
                p_opt.paragraph_format.space_after = Pt(1)
                r_opt = p_opt.add_run(f"□ {opt_letter}. ")
                r_opt.font.bold = True
                p_opt.add_run(opt_text)
            p_lbl = doc.add_paragraph()
            p_lbl.paragraph_format.left_indent = Inches(0.3)
            p_lbl.add_run("[Pilihan Ganda Kompleks: Jawaban benar lebih dari satu]").font.italic = True
        elif tipe == "Kategori":
            table = doc.add_table(rows=len(item["options"]) + 1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            hdr[0].text = "Pernyataan"
            hdr[1].text = "Benar"
            hdr[2].text = "Salah"
            for c in hdr:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                c.paragraphs[0].runs[0].font.bold = True
            for opt_idx, stmt in enumerate(item["options"]):
                row = table.rows[opt_idx + 1].cells
                row[0].text = f"{chr(65+opt_idx)}. {stmt}"
                row[1].text = "[   ]"
                row[2].text = "[   ]"
                row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                row[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_lbl = doc.add_paragraph()
            p_lbl.add_run("[Pilihan Ganda Kompleks: Kategori Benar / Salah]").font.italic = True
        else:
            for opt_idx, opt_text in enumerate(item["options"]):
                opt_letter = chr(65 + opt_idx)
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Inches(0.3)
                p_opt.paragraph_format.space_before = Pt(0)
                p_opt.paragraph_format.space_after = Pt(1)
                r_opt = p_opt.add_run(f"{opt_letter}. ")
                r_opt.font.bold = True
                p_opt.add_run(opt_text)
                
    doc.add_page_break()
    
    h_sec_b = doc.add_heading(level=2)
    h_sec_b.paragraph_format.space_before = Pt(8)
    h_sec_b.paragraph_format.space_after = Pt(8)
    r_b = h_sec_b.add_run("BAGIAN B: KUNCI JAWABAN & PEMBAHASAN LENGKAP")
    r_b.font.size = Pt(13)
    r_b.font.bold = True
    r_b.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)
    
    for item in paket_data:
        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.space_before = Pt(6)
        p_exp.paragraph_format.space_after = Pt(4)
        
        corr = item["correct"]
        if isinstance(corr, list):
            corr = ", ".join(corr)
        elif isinstance(corr, dict):
            corr = ", ".join([f"{k}: {v}" for k, v in corr.items()])
            
        r_num = p_exp.add_run(f"Soal Nomor {item['number']}  |  Kunci Jawaban: {corr}\n")
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        
        r_mat = p_exp.add_run(f"Materi / Lingkup Kompetensi: {item['materi']}\n")
        r_mat.font.italic = True
        r_mat.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        
        p_exp.add_run(f"Pembahasan & Langkah Penyelesaian:\n{item['explanation']}\n")
        
        p_line = doc.add_paragraph()
        p_line.add_run("―" * 45).font.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
        
    out_path = os.path.join(TARGET_DIR, filename)
    doc.save(out_path)
    print(f"[OK] Word Document: {out_path}")
    return out_path

def buat_quizizz_excel(paket_data, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quizizz Questions"
    
    headers = [
        "Question Text", "Question Type", "Option 1", "Option 2", "Option 3",
        "Option 4", "Option 5", "Correct Answer", "Time in seconds", "Image Link", "Explanation"
    ]
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        
    for row_idx, item in enumerate(paket_data, 2):
        q_text = item["question"].strip()
        opts = item["options"]
        tipe = item.get("tipe", "Pilihan Ganda")
        
        q_type_str = "Multiple Choice"
        corr_val = "1"
        if tipe == "MCMA":
            q_type_str = "Checkbox"
            corr_list = item["correct"] if isinstance(item["correct"], list) else ["A"]
            corr_val = ", ".join(str(ord(c) - ord('A') + 1) for c in corr_list)
        elif tipe == "Kategori":
            q_type_str = "Multiple Choice"
            corr_val = "1"
        else:
            c = str(item["correct"]).strip().upper()
            corr_val = ord(c) - ord('A') + 1 if len(c) == 1 and 'A' <= c <= 'D' else 1
            
        ws.cell(row=row_idx, column=1, value=q_text).alignment = align_left
        ws.cell(row=row_idx, column=2, value=q_type_str).alignment = align_center
        ws.cell(row=row_idx, column=3, value=opts[0] if len(opts) > 0 else "").alignment = align_left
        ws.cell(row=row_idx, column=4, value=opts[1] if len(opts) > 1 else "").alignment = align_left
        ws.cell(row=row_idx, column=5, value=opts[2] if len(opts) > 2 else "").alignment = align_left
        ws.cell(row=row_idx, column=6, value=opts[3] if len(opts) > 3 else "").alignment = align_left
        ws.cell(row=row_idx, column=7, value="").alignment = align_left
        ws.cell(row=row_idx, column=8, value=corr_val).alignment = align_center
        ws.cell(row=row_idx, column=9, value=60).alignment = align_center
        ws.cell(row=row_idx, column=10, value="").alignment = align_center
        ws.cell(row=row_idx, column=11, value=item["explanation"]).alignment = align_left
        
    out_path = os.path.join(TARGET_DIR, filename)
    try:
        wb.save(out_path)
        print(f"[OK] Quizizz Excel: {out_path}")
    except PermissionError:
        alt_path = os.path.join(TARGET_DIR, filename.replace('.xlsx', '_terbaru.xlsx'))
        wb.save(alt_path)
        print(f"[WARNING] {filename} sedang dibuka di Excel. Tersimpan sementara sebagai: {alt_path}")
    return out_path

def buat_pdf(paket_data, subject_name, paket_name, filename):
    out_path = os.path.join(TARGET_DIR, filename)
    doc = SimpleDocTemplate(out_path, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    style_normal = ParagraphStyle('CustomNormal', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=13, textColor=colors.HexColor('#1F2937'))
    style_heading = ParagraphStyle('CustomHeading', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11, leading=14, textColor=colors.HexColor('#1E3A8A'), spaceBefore=10, spaceAfter=6)
    
    story = []
    
    header_data = [[
        f"ASESMEN STANDARISASI PENDIDIKAN DAERAH / TKA SD 2026\n"
        f"TRY OUT DIGITAL INTERAKTIF: QUIZ WEB\n"
        f"{subject_name.upper()} - {paket_name.upper()}\n"
        f"Disusun Berdasarkan Kisi-Kisi Kurikulum Merdeka & Standar Try Out Resmi Pemerintah"
    ]]
    header_table = Table(header_data, colWidths=[490])
    header_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EBF5FB')),
        ('BOX', (0,0), (-1,-1), 1.5, colors.HexColor('#4472C4')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#4472C4')),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    story.append(Paragraph("BAGIAN A: NASKAH SOAL UJIAN (20 BUTIR)", style_heading))
    
    for item in paket_data:
        tipe = item.get("tipe", "Pilihan Ganda")
        q_text = f"<b>{item['number']}.</b> {item['question'].replace(chr(10), '<br/>')}"
        story.append(Paragraph(q_text, style_normal))
        story.append(Spacer(1, 4))
        
        if tipe == "MCMA":
            for opt_idx, opt_text in enumerate(item["options"]):
                opt_letter = chr(65 + opt_idx)
                story.append(Paragraph(f"□ <b>{opt_letter}.</b> {opt_text}", style_normal))
            story.append(Paragraph("<i>[Pilihan Ganda Kompleks: Jawaban benar lebih dari satu]</i>", style_normal))
        elif tipe == "Kategori":
            t_data = [["Pernyataan", "Benar", "Salah"]]
            for opt_idx, stmt in enumerate(item["options"]):
                t_data.append([f"{chr(65+opt_idx)}. {stmt}", "[   ]", "[   ]"])
            t = Table(t_data, colWidths=[370, 60, 60])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EBF5FB')),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
                ('ALIGN', (1,1), (-1,-1), 'CENTER'),
                ('FONTSIZE', (0,0), (-1,-1), 8.5),
            ]))
            story.append(t)
            story.append(Paragraph("<i>[Pilihan Ganda Kompleks: Kategori Benar / Salah]</i>", style_normal))
        else:
            for opt_idx, opt_text in enumerate(item["options"]):
                opt_letter = chr(65 + opt_idx)
                story.append(Paragraph(f"<b>{opt_letter}.</b> {opt_text}", style_normal))
        story.append(Spacer(1, 6))
        
    story.append(PageBreak())
    story.append(Paragraph("BAGIAN B: KUNCI JAWABAN & PEMBAHASAN LENGKAP", style_heading))
    
    for item in paket_data:
        corr = item["correct"]
        if isinstance(corr, list):
            corr = ", ".join(corr)
        elif isinstance(corr, dict):
            corr = ", ".join([f"{k}: {v}" for k, v in corr.items()])
        story.append(Paragraph(f"<b>Soal Nomor {item['number']}  |  Kunci Jawaban: {corr}</b>", style_normal))
        story.append(Paragraph(f"<i>Materi: {item['materi']}</i>", style_normal))
        story.append(Paragraph(f"Pembahasan:<br/>{item['explanation'].replace(chr(10), '<br/>')}", style_normal))
        story.append(Spacer(1, 6))
        
    doc.build(story)
    print(f"[OK] PDF: {out_path}")
    return out_path

# Execute for all 4 sets
tasks = [
    ("Matematika", "Paket 1", "Quiz_Web_Matematika_Paket_1"),
    ("Matematika", "Paket 2", "Quiz_Web_Matematika_Paket_2"),
    ("Bahasa Indonesia", "Paket 1", "Quiz_Web_Bahasa_Indonesia_Paket_1"),
    ("Bahasa Indonesia", "Paket 2", "Quiz_Web_Bahasa_Indonesia_Paket_2"),
]

for subj, pkt, base_name in tasks:
    data = grouped.get((subj, pkt), [])
    if data:
        print(f"\n=== MEMPROSES {base_name} ({len(data)} butir) ===")
        buat_docx(data, subj, pkt, f"{base_name}.docx")
        buat_quizizz_excel(data, f"{base_name}.xlsx")
        buat_pdf(data, subj, pkt, f"{base_name}.pdf")

print("\n[SUCCESS] Semua 12 file Quiz Web berhasil diperbarui dengan 3 model soal!")
