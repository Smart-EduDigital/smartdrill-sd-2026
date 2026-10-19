import openpyxl, os, json

folder = "C:/Users/MIP-USER/Documents/Project Soal"
xlsx_files = [f for f in os.listdir(folder) if f.endswith('.xlsx') and f.startswith('Quiz_Web')]
info = {}
for f in xlsx_files:
    path = os.path.join(folder, f)
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    rows = ws.max_row - 1
    info[f] = {'rows': rows, 'cols': ws.max_column}
print(json.dumps(info, indent=2))
